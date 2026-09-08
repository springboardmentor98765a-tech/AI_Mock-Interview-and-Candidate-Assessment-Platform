"""
Coding Practice — standalone from the scored live-interview flow.

Role locks language (see app/coding_bank.py); a candidate submits code,
it's run against a curated problem's stdin/stdout test cases by
app/judge.py, and scored purely algorithmically (exact output match) —
never by an AI judging the code. Every graded attempt is saved to
coding_submissions so the candidate can review their history and
export it as a PDF or Excel score list.
"""
import io
import json
import uuid
from datetime import datetime
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app import ai_providers
from app.coding_bank import (
    LANGUAGE_DISPLAY_NAME,
    ROLE_LANGUAGE_OPTIONS,
    get_problem,
    get_random_bank_problem,
    language_for_role,
    list_problems_for_language,
)
from app.database import get_db
from app.judge import run_submission
from app.models import CodingSubmission, GeneratedCodingQuestion
from app.schemas import (
    CodingQuestionOut,
    CodingSubmissionOut,
    GenerateCodingQuestionRequest,
    GeneratedCodingQuestionOut,
    RoleLanguageOut,
    SubmitCodeRequest,
    SubmitCodeResponse,
    TestCaseResultOut,
)
from app.security import CurrentUser, get_current_user, require_roles

router = APIRouter(prefix="/api/coding", tags=["coding-practice"])

VALID_DIFFICULTIES = {"easy", "medium", "hard"}


@router.get("/roles", response_model=list[RoleLanguageOut])
def list_roles():
    """Single source of truth for the role->language lock — the frontend
    renders its dropdown straight from this instead of hard-coding a
    second copy that could drift out of sync."""
    return ROLE_LANGUAGE_OPTIONS


@router.post("/generate", response_model=GeneratedCodingQuestionOut, status_code=201)
def generate_question(
    body: GenerateCodingQuestionRequest,
    db: Session = Depends(get_db),
    user: CurrentUser = Depends(require_roles("candidate")),
):
    """Generates ONE fresh coding problem with Gemini, locked to the
    language of `role`, at the requested difficulty. Gemini's answer key
    (reference solution) is run through the same judge a candidate's
    submission will go through before this ever returns — a hallucinated/
    inconsistent problem is silently discarded and regenerated, so a
    candidate is never handed broken test cases. If Gemini is
    unavailable/unconfigured or every attempt fails validation, falls
    back to a random problem from the curated offline bank so the
    feature always works."""
    language = language_for_role(body.role)
    difficulty = (body.difficulty or "medium").strip().lower()
    if difficulty not in VALID_DIFFICULTIES:
        difficulty = "medium"

    generated = ai_providers.generate_coding_question_llm(
        role=body.role,
        language=language,
        language_display_name=LANGUAGE_DISPLAY_NAME.get(language, language),
        difficulty=difficulty,
    )

    if generated:
        source = "ai"
        title = generated["title"]
        prompt_text = generated["prompt"]
        starter_code = generated["starter_code"]
        test_cases = generated["test_cases"]
    else:
        fallback = get_random_bank_problem(language, difficulty)
        source = "bank"
        title = fallback["title"]
        prompt_text = fallback["prompt"]
        starter_code = fallback["starter_code"]
        test_cases = fallback["test_cases"]

    question_id = uuid.uuid4().hex
    row = GeneratedCodingQuestion(
        id=question_id,
        candidate_id=user.id,
        role=body.role,
        language=language,
        difficulty=difficulty,
        title=title,
        prompt=prompt_text,
        starter_code=starter_code,
        test_cases_json=json.dumps(test_cases),
        source=source,
    )
    db.add(row)
    db.commit()

    return GeneratedCodingQuestionOut(
        id=question_id,
        title=title,
        difficulty=difficulty,
        prompt=prompt_text,
        language=language,
        starter_code=starter_code,
        source=source,
    )


@router.get("/questions", response_model=list[CodingQuestionOut])
def list_questions(
    role: str = Query(..., description="One of the role names from GET /roles"),
    user: CurrentUser = Depends(require_roles("candidate")),
):
    """Browse the curated offline bank directly (bypassing generation) —
    kept for reference/fallback use; the main practice flow uses
    POST /generate for a fresh AI question instead."""
    language = language_for_role(role)
    return list_problems_for_language(language)


@router.post("/submit", response_model=SubmitCodeResponse, status_code=201)
def submit_code(
    body: SubmitCodeRequest,
    db: Session = Depends(get_db),
    user: CurrentUser = Depends(require_roles("candidate")),
):
    language = language_for_role(body.role)

    # A questionId is either one this candidate got from POST /generate
    # (AI or offline-fallback — either way, its test cases are already
    # persisted), or a static id from GET /questions.
    generated = (
        db.query(GeneratedCodingQuestion)
        .filter(
            GeneratedCodingQuestion.id == body.questionId,
            GeneratedCodingQuestion.candidate_id == user.id,
        )
        .first()
    )

    if generated:
        title = generated.title
        test_cases = json.loads(generated.test_cases_json)
        language = generated.language
    else:
        problem = get_problem(body.questionId)
        if problem is None:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "Unknown coding question")
        title = problem["title"]
        test_cases = problem["test_cases"]

    if not body.code or not body.code.strip():
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Submitted code is empty")
    if len(body.code) > 50_000:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Submitted code is too long")

    result = run_submission(language, body.code, test_cases)

    submission = CodingSubmission(
        candidate_id=user.id,
        question_id=body.questionId,
        title=title,
        role=body.role,
        language=language,
        code=body.code,
        passed_count=result.passed_count,
        total_count=result.total_count,
        score_percent=result.score_percent,
    )
    db.add(submission)
    db.commit()
    db.refresh(submission)

    return SubmitCodeResponse(
        id=submission.id,
        questionId=body.questionId,
        title=title,
        language=language,
        compiled=result.compiled,
        compileError=result.compile_error,
        passedCount=result.passed_count,
        totalCount=result.total_count,
        scorePercent=result.score_percent,
        results=[
            TestCaseResultOut(
                input=r.input,
                expected_output=r.expected_output,
                actual_output=r.actual_output,
                passed=r.passed,
                error=r.error,
            )
            for r in result.results
        ],
    )


@router.get("/submissions/me", response_model=list[CodingSubmissionOut])
def my_submissions(
    db: Session = Depends(get_db),
    user: CurrentUser = Depends(require_roles("candidate")),
):
    rows = (
        db.query(CodingSubmission)
        .filter(CodingSubmission.candidate_id == user.id)
        .order_by(CodingSubmission.created_at.desc())
        .all()
    )
    return [
        CodingSubmissionOut(
            id=r.id,
            questionId=r.question_id,
            title=r.title,
            role=r.role,
            language=r.language,
            passedCount=r.passed_count,
            totalCount=r.total_count,
            scorePercent=r.score_percent,
            createdAt=r.created_at,
        )
        for r in rows
    ]


def _get_own_submissions(db: Session, user: CurrentUser):
    return (
        db.query(CodingSubmission)
        .filter(CodingSubmission.candidate_id == user.id)
        .order_by(CodingSubmission.created_at.desc())
        .all()
    )


@router.get("/submissions/export")
def export_submissions(
    format: str = Query("pdf", pattern="^(pdf|xlsx)$"),
    db: Session = Depends(get_db),
    user: CurrentUser = Depends(require_roles("candidate")),
):
    rows = _get_own_submissions(db, user)

    if format == "xlsx":
        buffer = _build_xlsx(rows, user)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        filename = "coding_practice_scores.xlsx"
    else:
        buffer = _build_pdf(rows, user)
        media_type = "application/pdf"
        filename = "coding_practice_scores.pdf"

    return StreamingResponse(
        buffer,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _build_xlsx(rows, user: CurrentUser) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Coding Practice Scores"

    ws.append([f"Coding Practice — Score List for {user.full_name or user.email}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}"])
    ws.append([])

    headers = ["Date", "Problem", "Role", "Language", "Passed", "Total", "Score %"]
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    for r in rows:
        ws.append(
            [
                r.created_at.strftime("%Y-%m-%d %H:%M") if r.created_at else "",
                r.title,
                r.role,
                LANGUAGE_DISPLAY_NAME.get(r.language, r.language),
                r.passed_count,
                r.total_count,
                r.score_percent,
            ]
        )

    for col, width in zip("ABCDEFG", (18, 26, 22, 20, 10, 10, 10)):
        ws.column_dimensions[col].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _build_pdf(rows, user: CurrentUser) -> io.BytesIO:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, title="Coding Practice Scores")
    styles = getSampleStyleSheet()
    elements = [
        Paragraph(f"Coding Practice — Score List", styles["Title"]),
        Paragraph(f"Candidate: {user.full_name or user.email}", styles["Normal"]),
        Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles["Normal"]),
        Spacer(1, 16),
    ]

    data = [["Date", "Problem", "Role", "Language", "Passed", "Total", "Score %"]]
    for r in rows:
        data.append(
            [
                r.created_at.strftime("%Y-%m-%d %H:%M") if r.created_at else "",
                r.title,
                r.role,
                LANGUAGE_DISPLAY_NAME.get(r.language, r.language),
                str(r.passed_count),
                str(r.total_count),
                f"{r.score_percent}%",
            ]
        )
    if len(data) == 1:
        data.append(["—", "No submissions yet", "—", "—", "—", "—", "—"])

    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2a44")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f5f9")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    buffer.seek(0)
    return buffer
